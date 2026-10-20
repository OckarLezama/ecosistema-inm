# -*- coding: utf-8 -*-
"""
IA-NAMI · LECTURA COMÚN DE LAS BASES
====================================
Lo usan todos los robots. Una base puede llegar en varias hojas de un mismo
archivo (cuando pasa del millón de renglones que caben en una hoja de Excel)
o en varios archivos (por año, por semana o por día). Aquí se juntan.

Reglas
  · Cada base se reconoce por sus columnas, no por el nombre del archivo.
  · Solo se leen hojas visibles; una hoja se junta si trae las mismas columnas
    que la primera (las tablas dinámicas o copias viejas se ignoran y se avisa).
  · Encabezados: mayúsculas, sin acentos ni espacios sobrantes; los errores de
    escritura conocidos se corrigen y se avisan (ALIAS).
  · Gana la entrega más reciente y cada archivo manda en todo su rango de fechas;
    las correcciones se reconocen por el nombre («CORRECCION 2026-09»). Ver [3].
  · Todo se toma desde INICIO (el periodo de la presidenta); lo anterior se avisa.
  · Nombre sugerido de los archivos: «FUENTE BASE AAAA-MM-DD a AAAA-MM-DD.xlsx».
    El nombre no decide nada: si sus fechas no coinciden con el contenido, se avisa.

ÍNDICE  [1] UTILIDADES   [2] HOJAS   [3] JUNTAR   [4] AVISOS   [5] BITÁCORA
"""
import os
import re
import unicodedata

import openpyxl
import pandas as pd

INICIO = pd.Timestamp("2024-10-01")
ALIAS = {"NACIONADLIDAD": "NACIONALIDAD", "NACIONALIDA": "NACIONALIDAD"}
RENOMBRES = [("FECHA", "DIA", {"LOCATION", "CIUDAD EEUU", "TOTAL"}),                      # CBP desde oct 2026
           ("ENCUENTRO", "AGENCIA", {"LOCATION", "CIUDAD EEUU", "TOTAL"}),
           ("FECHA_REPATRIACION", "DIA", {"PUNTO_REPATRIACION", "MEXICANOS REPATRIADOS"}),   # DGPMV desde oct 2026
           ("ESTADO", "O.R.", {"PUNTO_REPATRIACION", "MEXICANOS REPATRIADOS"})]
AVISOS = {}          # nombre de la base → observaciones (sin repetir)
PERIODOS = {}        # nombre de la base → {desde, hasta, archivo, filas}: lo que lee cada robot fuera de datos.js (página Bases)


# ==== [1] UTILIDADES ==============================================
def norm(texto):
    t = unicodedata.normalize("NFD", str(texto).upper())
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn")
    return " ".join(t.split())


def periodo(base, desde, hasta, archivo, filas=None):
    PERIODOS[base] = {"desde": str(desde)[:10], "hasta": str(hasta)[:10], "archivo": archivo, "filas": None if filas is None else int(filas)}


def avisar(base, texto):
    l = AVISOS.setdefault(base, [])
    if texto not in l:
        l.append(texto)


def fechas_del_nombre(nombre):
    """«… 2026-10-13 a 2026-10-19.xlsx» → (desde, hasta); «… al 2026-10-12» o «… 2025» también."""
    nombre = sin_entrega(nombre)
    f = re.findall(r"(20\d\d)-(\d\d)-(\d\d)", nombre)
    if f:
        d = [pd.Timestamp(int(a), int(m), int(x)) for a, m, x in f]
        return d[0], d[-1]
    a = re.findall(r"\b(20\d\d)\b", nombre)
    if len(a) == 1:
        return int(a[0])                                    # solo el año: basta con que el contenido caiga en ese año
    return None


# ==== [2] HOJAS ===================================================
def encabezados(fila, base=None):
    cols = []
    for c in fila:
        n = norm(c)
        if n in ALIAS:
            if base:
                avisar(base, f"Encabezado «{c}» leído como «{ALIAS[n]}» (corregir en el archivo).")
            n = ALIAS[n]
        cols.append(n)
    # Encabezados que cambiaron de nombre en entregas recientes: se leen con el nombre de siempre solo si la base
    # trae sus columnas propias y le falta la de siempre (así no se confunde «FECHA» o «ESTADO» de otra base).
    for viejo, nuevo, si in RENOMBRES:
        if viejo in cols and nuevo not in cols and si <= set(cols):
            if base:
                avisar(base, f"Encabezado «{viejo}» leído como «{nuevo}».")
            cols[cols.index(viejo)] = nuevo
    return cols


def hojas(ruta):
    """[(nombre de hoja, tabla sin encabezado)] de las hojas visibles, y la lista de ocultas."""
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    vis = [h.title for h in libro.worksheets if h.sheet_state == "visible"]
    ocu = [h.title for h in libro.worksheets if h.sheet_state != "visible"]
    libro.close()
    return [(h, pd.read_excel(ruta, sheet_name=h, header=None)) for h in vis], ocu


def primera_fila(ruta, acepta=None):
    """Encabezados de la primera hoja visible, sin leer todo el archivo. Con `acepta` (función sobre la lista de
    encabezados), la primera hoja visible que la cumple: así una hoja de apoyo puesta antes que la base no estorba."""
    libro = openpyxl.load_workbook(ruta, read_only=True)
    fila = ()
    for h in libro.worksheets:
        if h.sheet_state != "visible":
            continue
        f = [c for c in next(h.iter_rows(max_row=1, values_only=True), ()) if c is not None]
        if acepta is None or acepta(f):
            fila = f
            break
    libro.close()
    return list(fila)


def tabla_de(ruta, base=None, acepta=None):
    """Junta las hojas visibles que traen las mismas columnas que la primera (o que la primera que cumple `acepta`).
    Devuelve una tabla con encabezados normalizados y la lista de hojas que no se leyeron."""
    hs, ocultas = hojas(ruta)
    if acepta is not None:
        k = next((i for i, (_, t) in enumerate(hs) if len(t) and acepta([c for c in t.iloc[0].tolist() if pd.notna(c)])), None)
        if k:
            hs = hs[k:] + hs[:k]
    if not hs:
        return None, ocultas
    orig = [" ".join(str(c).split()) for c in hs[0][1].iloc[0].tolist()]
    cols = encabezados(hs[0][1].iloc[0].tolist(), base)
    partes, fuera = [], list(ocultas)
    for nombre, t in hs:
        c = encabezados(t.iloc[0].tolist(), base)
        if set(c) != set(cols):
            fuera.append(nombre)
            continue
        x = t.iloc[1:].copy()
        x.columns = c
        partes.append(x[cols])
    t = pd.concat(partes, ignore_index=True)
    t.attrs["orig"] = orig                                   # encabezados tal como vienen (para mostrar nombres)
    return t, fuera


# ==== [3] JUNTAR ==================================================
# Reglas de carga (también para correcciones de un mes completo):
#   1. Gana la entrega más reciente. Fecha de entrega: la que diga el nombre («… entregado 2026-10-17»);
#      si no dice, cuándo entró el archivo al repositorio (git). Si no se sabe, una corrección va después de lo que
#      corrige y, entre envíos normales, el que llega más lejos (como antes).
#   2. Cada archivo manda en todo su rango de fechas: los días de ese rango se toman solo de él, aunque un
#      día no venga (cuenta como cero) y aunque una entrega anterior sí lo trajera. Nada se suma dos veces.
#   3. Tipo de envío por el nombre: si dice «CORRECCION» es una corrección y su rango es el que dice el nombre
#      («… CORRECCION 2026-09.xlsx» = todo septiembre; «… CORRECCION 2026-09-01 a 2026-09-15.xlsx»);
#      cualquier otro es un envío de días nuevos y su rango es del primer al último día que trae.
#   4. Cada corrida deja huella por base y día (renglones y total); la siguiente compara y anota en la bitácora
#      qué cambió (datos/bitacora.js, sección Bases).
def dia(serie):
    return pd.to_datetime(serie, errors="coerce", dayfirst=True).dt.normalize()


def es_correccion(nombre):
    return "CORRECCION" in norm(nombre)


def rango_del_nombre(nombre):
    """Rango que declara el nombre: «2026-09-01 a 2026-09-30», «2026-09» (el mes completo) o «2026-08 a 2026-09»."""
    nombre = sin_entrega(nombre)
    f = re.findall(r"(20\d\d)-(\d\d)-(\d\d)", nombre)
    if f:
        d = [pd.Timestamp(int(a), int(m), int(x)) for a, m, x in f]
        return min(d), max(d)
    m = re.findall(r"(20\d\d)-(\d\d)(?!-?\d)", nombre)
    if m:
        d = [pd.Timestamp(int(a), int(b), 1) for a, b in m]
        return min(d), max(d) + pd.offsets.MonthEnd(0)
    return None


def sin_entrega(nombre):
    """El nombre sin «entregado AAAA-MM-DD» (esa fecha es de la entrega, no del periodo que trae)."""
    return re.sub(r"(?i)\b(entregad[oa]|recibid[oa]|enviad[oa]|entrega)\s+20\d\d-\d\d-\d\d", "", nombre)


FECHA_ENTREGA = re.compile(r"(?:ENTREGAD[OA]|RECIBID[OA]|ENVIAD[OA]|ENTREGA)\s+(20\d\d)-(\d\d)-(\d\d)")


def entrega_de(nombre, origen=None):
    """Momento de la entrega (segundos): lo que diga el nombre, si no cuándo entró al repositorio (git); si no, 0."""
    m = FECHA_ENTREGA.search(norm(nombre))
    if m:
        return pd.Timestamp(int(m.group(1)), int(m.group(2)), int(m.group(3)), 23, 59).timestamp(), "nombre"
    if origen:
        ruta = os.path.join(origen, nombre)
        try:
            import subprocess
            r = subprocess.run(["git", "log", "--diff-filter=AM", "-1", "--format=%ct", "--", nombre], cwd=origen, capture_output=True, text=True, timeout=20)
            if r.returncode == 0 and r.stdout.strip():
                return float(r.stdout.strip()), "repositorio"
        except Exception:
            pass
    return 0.0, "sin fecha"      # la fecha del archivo en disco no sirve (al copiar o clonar cambia): se ordena por tipo y alcance


def por_entrega(rutas, origen):
    """Rutas ordenadas de la entrega más antigua a la más reciente (para bases que llegan completas cada vez: cortes y validaciones)."""
    return sorted(rutas, key=lambda r: (entrega_de(os.path.basename(r), origen)[0], os.path.basename(r)))


MEDIDAS = ("TOTAL", "PRESENTADOS", "RETORNADOS A SU PAIS", "EXTRANJEROS RECIBIDOS DE EE.UU.", "MEXICANOS REPATRIADOS", "PERSONAS")
HUELLAS = {}          # base → {día: [renglones, total]} de esta corrida (para la bitácora)
CAMBIOS = {}          # base → [archivos más recientes y su tipo] (para la bitácora)


def huella(base, t, col):
    med = next((c for c in MEDIDAS if c in t.columns), None)
    g = t.groupby(col)
    filas = g.size()
    tot = g[med].apply(lambda s: float(pd.to_numeric(s, errors="coerce").fillna(0).sum())) if med else filas.astype(float)
    HUELLAS[base] = {str(d.date()): [int(filas[d]), round(float(tot[d]), 2)] for d in filas.index}


def juntar(partes, base, col_dia="DIA", origen=None):
    """partes: [(archivo, tabla con encabezados normalizados)]. Junta según las reglas de arriba y recorta antes de INICIO."""
    if col_dia not in partes[0][1].columns:                 # sin fechas (cortes): la entrega más reciente
        if len(partes) > 1:
            k = max(range(len(partes)), key=lambda i: (entrega_de(partes[i][0], origen)[0], partes[i][0]))
            avisar(base, "Hay " + str(len(partes)) + " archivos de esta base y no trae fechas: se usa la entrega más reciente, " + partes[k][0] + ".")
            return partes[k][1].reset_index(drop=True)
        return partes[-1][1].reset_index(drop=True)
    info = []
    for nombre, t in partes:
        t = t.copy()
        t["_D"] = dia(t[col_dia])
        sin = int(t["_D"].isna().sum())
        if sin:
            avisar(base, f"{nombre}: {sin:,} renglones no traen una fecha válida y no se cuentan.")
        t = t[t["_D"].notna()]
        if not len(t):
            avisar(base, f"{nombre}: no trae renglones con fecha.")
            continue
        a, b = t["_D"].min(), t["_D"].max()
        corr = es_correccion(nombre)
        fn = fechas_del_nombre(nombre)
        rn = rango_del_nombre(nombre) if corr else None
        if corr and rn is None:
            avisar(base, f"{nombre}: dice CORRECCION pero el nombre no trae el periodo (ej. «… CORRECCION 2026-09.xlsx»); se corrige del {a.date()} al {b.date()}, que es lo que trae.")
        if rn:
            fuera = (t["_D"] < rn[0]) | (t["_D"] > rn[1])
            if fuera.any():
                avisar(base, f"{nombre}: es corrección del {rn[0].date()} al {rn[1].date()} y trae {int(fuera.sum()):,} renglones fuera de ese periodo; no se usan.")
                t = t[~fuera]
            r0, r1 = rn
        else:
            r0, r1 = a, b
            if isinstance(fn, int) and (a.year != fn or b.year != fn):
                avisar(base, f"{nombre}: el nombre dice {fn}, pero trae {a.date()} a {b.date()}.")
            elif isinstance(fn, tuple) and (fn[0] > a + pd.Timedelta(days=6) or fn[1] < b - pd.Timedelta(days=6) or fn[1] > b + pd.Timedelta(days=6)):
                avisar(base, f"{nombre}: el nombre dice {fn[0].date()} a {fn[1].date()}, pero trae {a.date()} a {b.date()}.")
        ent, fuente = entrega_de(nombre, origen)
        info.append({"n": nombre, "t": t, "r0": r0, "r1": r1, "b": b, "corr": corr, "ent": ent, "fte": fuente})
    if not info:
        return partes[0][1].iloc[0:0]
    # gana la entrega más reciente; si no se sabe, una corrección va después de lo que corrige y luego el que llega más lejos
    info.sort(key=lambda x: (x["ent"], x["corr"], x["b"], x["n"]))
    tomados = []          # rangos ya cubiertos por entregas más recientes
    salida, filas = [], {}
    for x in reversed(info):
        t, d = x["t"], x["t"]["_D"]
        tapado = pd.Series(False, index=t.index)
        for (q0, q1, quien) in tomados:
            tapado |= (d >= q0) & (d <= q1)
        if tapado.any():
            rep = sorted(set(d[tapado].unique()))
            n_viejo = t[tapado].groupby("_D").size()
            bajos = [(dd.date(), int(n_viejo[dd]), filas.get(dd, 0)) for dd in rep if n_viejo[dd] >= 20 and filas.get(dd, 0) < 0.5 * n_viejo[dd]]
            nuevo_es_corr = any(q[2]["corr"] for q in tomados)
            if bajos and not nuevo_es_corr:
                avisar(base, "Revisar: " + str(len(bajos)) + " días llegan en la entrega más reciente con menos de la mitad de renglones que en " + x["n"] +
                       " (" + ", ".join(f"{a}: {nv:,} → {nn:,}" for a, nv, nn in bajos[:5]) + (" …" if len(bajos) > 5 else "") + "). Puede ser un envío incompleto.")
            quien = next(q[2]["n"] for q in tomados if ((d[tapado] >= q[0]) & (d[tapado] <= q[1])).any())
            avisar(base, f"{x['n']} es una entrega anterior: sus {len(rep):,} días del periodo que cubre {quien} se tomaron de ese archivo. No se suma nada dos veces.")
            vacios = [dd for dd in rep if dd not in filas]
            if vacios:
                avisar(base, f"{len(vacios):,} días que traía {x['n']} no vienen en la entrega más reciente que cubre ese periodo; se cuentan en cero"
                             f" ({', '.join(str(v.date()) for v in vacios[:5])}{' …' if len(vacios) > 5 else ''}). Si faltan por error, pedir el archivo completo.")
            t = t[~tapado]
        for dd, n in t.groupby("_D").size().items():
            filas.setdefault(dd, int(n))
        tomados.append((x["r0"], x["r1"], x))
        salida.append(t)
    CAMBIOS[base] = [{"archivo": x["n"], "tipo": "corrección" if x["corr"] else "días nuevos", "desde": str(x["r0"].date()), "hasta": str(x["r1"].date()),
                      "entrega": pd.Timestamp(x["ent"], unit="s").strftime("%Y-%m-%d") if x["ent"] else "", "segun": x["fte"]} for x in info]
    t = pd.concat(salida[::-1], ignore_index=True).sort_values("_D", kind="stable")
    antes = t["_D"] < INICIO
    if antes.any():
        avisar(base, f"{int(antes.sum()):,} renglones son de antes del 1 oct 2024 (el más antiguo, del {t['_D'].min().date()}). "
                     "No se usan: la plataforma empieza el 1 oct 2024.")
        t = t[~antes]
    t[col_dia] = t["_D"]
    huella(base, t, "_D")
    return t.drop(columns="_D").reset_index(drop=True)


def leer_por_columnas(origen, columnas, base, sin=frozenset()):
    """Busca en origen todos los archivos cuya primera hoja visible trae `columnas` (y ninguna de `sin`),
    junta hojas y archivos, y devuelve (tabla, [archivos])."""
    partes = []
    for ruta in sorted(os.path.join(origen, f) for f in os.listdir(origen) if f.lower().endswith(".xlsx") and not f.startswith("~$")):
        cols = set(encabezados(primera_fila(ruta)))
        if columnas <= cols and not (sin & cols):
            t, fuera = tabla_de(ruta, base)
            if fuera:
                avisar(base, os.path.basename(ruta) + ": las hojas «" + "», «".join(fuera) + "» no se usaron (están ocultas o no tienen las columnas de la base). Si traen datos que deben contar, pedir al área que los pase a la hoja principal.")
            partes.append((os.path.basename(ruta), t))
    if not partes:
        return None, []
    return juntar(partes, base, origen=origen), [p[0] for p in partes]


# ==== [3b] TRÁMITES: UN TRÁMITE, UN DOCUMENTO =======================
DOC = ("CONSTANCIA", "OFICIO", "PERMISO", "SIN EXPEDICION", "TARJETA")
SIN_TIPO = "TARJETA SIN TIPO REGISTRADO"


def documentos_tramite(t):
    """Desde 2026 la base de DGRAM trae un renglón por trámite y marca con 1 cada documento. Trae dos errores
    de captura que aquí se corrigen y se avisan: trámites con tarjeta que también marcan «Sin expedición»
    (no cuenta), y trámites sin ningún documento marcado (con tarjeta: «Tarjeta sin tipo registrado»;
    sin tarjeta: «Sin expedición»). En el formato anterior (renglones agrupados con conteos) no cambia nada."""
    docs = [c for c in t.columns if c.startswith(DOC)]
    sin = next((c for c in docs if c.startswith("SIN EXPEDICION")), None)
    X = t[docs].apply(pd.to_numeric, errors="coerce").fillna(0)
    if X.values.max(initial=0) > 1:                          # formato agrupado (conteos): se usa tal cual
        t[docs] = X
        return t, docs
    con = t["TIPO DE CONCLUSION"].map(norm).str.startswith("CON EXPEDICION") if "TIPO DE CONCLUSION" in t.columns else pd.Series(False, index=t.index)
    tarj = X[[c for c in docs if c.startswith("TARJETA")]].sum(axis=1) > 0
    if sin:
        doble = con & tarj & (X[sin] > 0)
        if doble.any():
            avisar("Trámites", f"{int(doble.sum()):,} trámites con tarjeta marcan también «Sin expedición de documento»; se cuentan solo con su tarjeta.")
            X.loc[doble, sin] = 0
    vacio = X.sum(axis=1) == 0
    X[SIN_TIPO] = 0
    if vacio.any():
        a, b = int((vacio & con).sum()), int((vacio & ~con).sum())
        avisar("Trámites", f"{int(vacio.sum()):,} trámites no marcan ningún documento: {a:,} terminaron con tarjeta (se cuentan como «Tarjeta sin tipo registrado») y {b:,} sin tarjeta (se cuentan como «Sin expedición»).")
        X.loc[vacio & con, SIN_TIPO] = 1
        if sin:
            X.loc[vacio & ~con, sin] = 1
    res = next((c for c in t.columns if c.startswith("TIPO DE RESOLUCION")), None)
    if res is not None:                                       # concluidos que la base aún marca «en proceso» o «cancelado»
        r = t[res].map(norm)
        ep, ca = r.str.startswith("EN PROCESO"), r.str.startswith("CANCELAD")
        if ep.any() or ca.any():
            avisar("Trámites", f"Entre los trámites concluidos, «Tipo de resolución» dice {int(ep.sum()):,} «En proceso» y {int(ca.sum()):,} «Cancelados»; "
                   f"{int((ep & con).sum()):,} de los «En proceso» ya dicen «Con expedición de tarjeta». Se cuentan por «Tipo de conclusión»; confirmar con el área.")
    mult = X.sum(axis=1) > 1
    if mult.any():
        avisar("Trámites", f"{int(mult.sum()):,} trámites marcan más de un documento; cuentan una vez por documento.")
    for c in X.columns:
        t[c] = X[c]
    return t, docs + [SIN_TIPO]


# ==== [4] AVISOS ==================================================
def avisos_de(base):
    return AVISOS.get(base, [])


def guardar(destino, robot):
    """Escribe las observaciones de este robot en datos/observaciones.js (cada robot reemplaza solo las suyas).
    La sección Bases de la página las muestra junto a cada base."""
    import json
    ruta = os.path.join(destino, "observaciones.js")
    todo = {}
    if os.path.exists(ruta):
        txt = open(ruta, encoding="utf-8").read()
        try:
            todo = json.loads(txt[txt.index("{"):txt.rindex("}") + 1])
        except ValueError:
            todo = {}
    todo[robot] = AVISOS
    bitacora(destino, robot)
    if PERIODOS:
        todo.setdefault("_periodos", {}).update(PERIODOS)
    os.makedirs(destino, exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as f:
        f.write("// Observaciones de cada base, por robot. Lo escribe robot/lector.py · no editar a mano\n")
        f.write("window.IANAMI_OBS=" + json.dumps(todo, ensure_ascii=False, separators=(",", ":")) + ";\n")


# ==== [5] BITÁCORA ================================================
def _carga_js(ruta, var):
    import json
    if not os.path.exists(ruta):
        return None
    txt = open(ruta, encoding="utf-8").read()
    try:
        return json.loads(txt[txt.index("[", txt.index(var)):txt.rindex("]") + 1])
    except ValueError:
        return None


def bitacora(destino, robot, hoy=None):
    """Compara la huella de cada base con la corrida anterior y anota lo que cambió: días con otra cifra, días que se
    quitaron y días nuevos, con la entrega responsable. datos/huella.json guarda la última huella; datos/bitacora.js,
    las anotaciones (lo más reciente primero). La primera corrida solo deja la huella."""
    import json
    if not HUELLAS:
        return
    hoy = hoy or pd.Timestamp.now().strftime("%Y-%m-%d %H:%M")
    rh = os.path.join(destino, "huella.json")
    prev = json.load(open(rh, encoding="utf-8")) if os.path.exists(rh) else {}
    rb = os.path.join(destino, "bitacora.js")
    lista = _carga_js(rb, "IANAMI_BITACORA") or []
    for base, h in HUELLAS.items():
        clave = robot + "|" + base
        p = prev.get(clave)
        prev[clave] = h
        if p is None:
            continue
        cambio = sorted(d for d in h if d in p and h[d] != p[d])
        quitados = sorted(d for d in p if d not in h)
        nuevos = sorted(d for d in h if d not in p)
        if not (cambio or quitados):
            if nuevos:
                lista.insert(0, {"fecha": hoy, "robot": robot, "base": base, "tipo": "días nuevos", "nuevos": len(nuevos), "desde": nuevos[0], "hasta": nuevos[-1],
                                 "total": round(sum(h[d][1] for d in nuevos), 2)})
            continue
        dias = cambio + quitados
        antes = round(sum(p[d][1] for d in dias), 2)
        despues = round(sum(h[d][1] for d in cambio), 2)
        porm = {}
        for d in dias:
            m = porm.setdefault(d[:7], [0, 0.0, 0.0])
            m[0] += 1; m[1] += p[d][1]; m[2] += h[d][1] if d in h else 0
        resp = [c for c in CAMBIOS.get(base, []) if c["desde"] <= dias[-1] and c["hasta"] >= dias[0]]
        resp = resp[-1:]                                    # la entrega más reciente que cubre esos días es la que los cambió
        lista.insert(0, {"fecha": hoy, "robot": robot, "base": base, "tipo": "corrección" if any(c["tipo"] == "corrección" for c in resp) else "cambio en días ya cargados",
                         "archivos": [c["archivo"] for c in resp], "dias": len(dias), "quitados": len(quitados), "desde": dias[0], "hasta": dias[-1],
                         "antes": antes, "despues": despues, "meses": [[k, v[0], round(v[1], 2), round(v[2], 2)] for k, v in sorted(porm.items())],
                         "detalle": [[d, p[d][1], h[d][1] if d in h else 0] for d in dias[:62]], "nuevos": len(nuevos)})
    os.makedirs(destino, exist_ok=True)
    json.dump(prev, open(rh, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    with open(rb, "w", encoding="utf-8") as f:
        f.write("// Bitácora de cambios en los datos: qué cambió en cada carga, por base y día. Lo escribe robot/lector.py · no editar a mano\n")
        f.write("window.IANAMI_BITACORA=" + json.dumps(lista[:400], ensure_ascii=False, separators=(",", ":")) + ";\n")
    HUELLAS.clear()
